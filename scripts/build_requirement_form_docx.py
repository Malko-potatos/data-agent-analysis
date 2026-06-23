from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "templates" / "data-analysis-agent-requirement-form.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE_GRAY = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "C9D3DF"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=INK, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def set_table_geometry(table, widths, indent_dxa=120, cell_margin=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for side, value in (("top", 80), ("bottom", 80), ("start", cell_margin), ("end", cell_margin)):
        child = cell_mar.find(qn(f"w:{side}"))
        if child is None:
            child = OxmlElement(f"w:{side}")
            cell_mar.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def mark_first_row_as_header(table):
    row = table.rows[0]
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_spacer(doc, pts=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("데이터 분석 요청 카드")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = doc.add_paragraph()
    p.style = "Subtitle"
    p.add_run("비개발자가 먼저 채우고, 에이전트가 작업 방식과 검증 기준을 정리하는 쉬운 양식")


def add_note_box(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360], indent_dxa=120)
    set_table_borders(table, color="D6DEE8")
    mark_first_row_as_header(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(10.5)
    title_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(body)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(INK)
    add_spacer(doc, 6)


def add_label_detail_table(doc, rows, widths=(1850, 7510), header=None):
    extra = 1 if header else 0
    table = doc.add_table(rows=len(rows) + extra, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, list(widths), indent_dxa=120)
    set_table_borders(table)
    mark_first_row_as_header(table)

    start = 0
    if header:
        cell = table.cell(0, 0).merge(table.cell(0, 1))
        set_cell_shading(cell, LIGHT_BLUE_GRAY)
        set_cell_text(cell, header, bold=True, color=DARK_BLUE, size=10.5)
        start = 1

    for idx, (label, guide) in enumerate(rows, start=start):
        left, right = table.cell(idx, 0), table.cell(idx, 1)
        set_cell_shading(left, LIGHT_GRAY)
        set_cell_text(left, label, bold=True, color=DARK_BLUE, size=10)
        set_cell_text(right, guide, color=MUTED, size=10)

    add_spacer(doc, 8)
    return table


def add_three_col_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE_GRAY):
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, widths, indent_dxa=120)
    set_table_borders(table)
    mark_first_row_as_header(table)
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, header_fill)
        set_cell_text(cell, header, bold=True, color=DARK_BLUE, size=10)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_text(cell, text, color=MUTED if c_idx > 0 else INK, size=9.8)
    add_spacer(doc, 8)
    return table


def add_checklist_table(doc, title, items):
    table = doc.add_table(rows=len(items) + 1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [520, 8840], indent_dxa=120)
    set_table_borders(table)
    mark_first_row_as_header(table)
    header = table.cell(0, 0).merge(table.cell(0, 1))
    set_cell_shading(header, LIGHT_BLUE_GRAY)
    set_cell_text(header, title, bold=True, color=DARK_BLUE, size=10.5)
    for idx, item in enumerate(items, start=1):
        set_cell_text(table.cell(idx, 0), "[ ]", color=DARK_BLUE, size=10)
        table.cell(idx, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(table.cell(idx, 1), item, color=INK, size=10)
    add_spacer(doc, 8)


def add_response_box(doc, title, guide, height_rows=4):
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360], indent_dxa=120)
    set_table_borders(table)
    mark_first_row_as_header(table)
    set_cell_shading(table.cell(0, 0), LIGHT_BLUE_GRAY)
    set_cell_text(table.cell(0, 0), title, bold=True, color=DARK_BLUE, size=10.5)
    cell = table.cell(1, 0)
    set_cell_text(cell, guide, color=MUTED, size=10)
    for _ in range(height_rows):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(" ")
    add_spacer(doc, 8)


def add_page_break(doc):
    doc.add_page_break()


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(10)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "Data Analysis Agent Workbook"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Calibri"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = "비개발자는 앞쪽의 쉬운 질문만 작성하고, 나머지는 에이전트와 함께 정리합니다."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.name = "Calibri"
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def build():
    doc = Document()
    configure_styles(doc)
    add_title(doc)
    add_note_box(
        doc,
        "작성 방법",
        "처음 작성자는 앞쪽의 쉬운 질문만 채우면 됩니다. use case, harness, 검증 방식 같은 어려운 분류는 "
        "에이전트가 이 내용을 보고 뒤쪽에 정리합니다.",
    )

    add_label_detail_table(
        doc,
        [
            ("요청 이름", "예: 6월 마케팅 성과 분석, 수업 진행률 대시보드"),
            ("작성자", "이름 또는 팀"),
            ("작성일", "YYYY-MM-DD"),
            ("함께 볼 사람", "예: 대표, 마케팅팀, 운영팀, 재무 담당자"),
        ],
        header="기본 정보",
    )

    add_response_box(
        doc,
        "1. 무엇을 알고 싶나요?",
        "예: 어떤 광고가 효과적인지 알고 싶다. 프로그램별 수업 진행률을 보고 싶다.",
        height_rows=3,
    )

    add_response_box(
        doc,
        "2. 이 결과로 어떤 결정을 하고 싶나요?",
        "예: 광고 예산을 어디에 더 쓸지 정한다. 지연 위험이 큰 수업을 먼저 챙긴다.",
        height_rows=3,
    )

    add_checklist_table(
        doc,
        "3. 데이터는 어디에 있나요? 아는 것만 체크하세요.",
        [
            "Excel, CSV, Google Sheet 같은 파일",
            "회사 DB 또는 어드민 화면",
            "GA, 광고 플랫폼, CRM 같은 외부 서비스",
            "이메일, Slack, 설문, 리뷰, 커뮤니티 글",
            "아직 잘 모르겠다",
        ],
    )

    add_response_box(
        doc,
        "데이터 위치를 조금 더 적어주세요.",
        "예: 파일명, 시트 이름, 어드민 메뉴, 대시보드 링크, 담당자 이름",
        height_rows=2,
    )

    add_checklist_table(
        doc,
        "4. 어떤 결과물이 있으면 좋나요?",
        [
            "한 번 보는 분석 리포트",
            "표나 차트",
            "반복해서 보는 대시보드",
            "매일/매주 자동 알림",
            "개선안 또는 다음 액션 목록",
        ],
    )

    add_spacer(doc, 10)
    add_response_box(
        doc,
        "5. 지금은 어떻게 하고 있나요?",
        "예: 어드민에서 내려받아 엑셀로 계산한다. 담당자가 매주 수동으로 정리한다.",
        height_rows=3,
    )

    add_checklist_table(
        doc,
        "6. 꼭 조심해야 할 일이 있나요?",
        [
            "개인정보나 민감정보가 있다.",
            "원본 파일이나 DB를 수정하면 안 된다.",
            "숫자가 틀리면 중요한 의사결정에 영향을 준다.",
            "외부로 메일, 메시지, 게시글을 보내면 안 된다.",
            "실행 전에 사람이 승인해야 한다.",
            "잘 모르겠다. 에이전트가 위험을 찾아줬으면 좋겠다.",
        ],
    )

    add_response_box(
        doc,
        "7. 맞았는지 무엇과 비교하면 좋을까요?",
        "예: 기존 어드민 숫자, 지난달 보고서, DB 총합, 담당자가 이미 알고 있는 기준 숫자",
        height_rows=3,
    )

    add_response_box(
        doc,
        "추가로 남기고 싶은 말",
        "불편한 점, 원하는 형식, 꼭 포함해야 할 항목, 걱정되는 점을 자유롭게 적습니다.",
        height_rows=3,
    )

    add_spacer(doc, 10)
    doc.add_heading("에이전트와 함께 정리하는 영역", level=1)
    add_note_box(
        doc,
        "에이전트 사용 안내",
        "아래 영역은 비개발자가 혼자 작성하지 않아도 됩니다. 위 내용을 에이전트에게 보여주고, "
        "에이전트가 빠진 질문, 작업 유형, 검증 기준, 실행 프롬프트를 채우게 합니다.",
    )

    add_label_detail_table(
        doc,
        [
            ("작업 유형", "분석 / 데이터 정리 / 대시보드 / 자동 알림 / 기존 agent 점검 중 무엇에 가까운가?"),
            ("추가 질문", "시작 전에 사용자에게 물어볼 질문 3개 이하"),
            ("필요한 자료", "파일, 링크, 계정, 권한, 기존 보고서 등"),
            ("조심할 경계", "수정 금지, 발송 금지, 개인정보, 승인 필요 항목"),
            ("맞춤 확인", "어떤 숫자나 자료와 비교해 결과를 확인할 것인가?"),
        ],
        header="에이전트 정리 카드",
    )

    add_response_box(
        doc,
        "에이전트에게 보낼 실행 요청 초안",
        "위 요청 카드를 바탕으로 필요한 자료를 먼저 확인하고, 분석/자동화 방식을 제안한 뒤, "
        "검증 기준과 승인 필요한 일을 분리해서 실행 계획을 작성해주세요.",
        height_rows=5,
    )

    add_label_detail_table(
        doc,
        [
            ("사용자 확인", "이해한 내용이 맞음 / 일부 수정 필요 / 더 질문 필요"),
            ("다음 단계", "자료 전달 / 에이전트 실행 / 대시보드 설계 / 자동화 보류"),
            ("승인자", "필요할 때만 작성"),
        ],
        header="실행 전 확인",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
